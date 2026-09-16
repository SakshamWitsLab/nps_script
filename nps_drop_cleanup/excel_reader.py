"""Parsing of the DROP sheets in the triage workbook.

Expected workbook layout (from ``NPS_RAG_Ingestion_Review_TRIAGED.xlsx``):

* ``DROP - URLs``      - column 'Example full URL' and 'Canonical page (...)'.
* ``DROP - Documents`` - column 'File name'.

Column names are configurable through env (``COL_*``).
"""

from dataclasses import dataclass, field

from openpyxl import load_workbook

from config import Config
from normalizer import normalize_url


class SheetFormatError(RuntimeError):
    pass


@dataclass
class UrlTarget:
    source_column: str
    original: str
    normalized: str
    section: str = ""

    @property
    def is_canonical(self) -> bool:
        return "canonical" in self.source_column.lower()


@dataclass
class FileTarget:
    original: str
    section: str = ""


@dataclass
class ExcelData:
    url_targets: list[UrlTarget] = field(default_factory=list)
    file_targets: list[FileTarget] = field(default_factory=list)
    url_rows_total: int = 0
    url_rows_skipped: int = 0
    file_rows_total: int = 0
    file_rows_skipped: int = 0

    @property
    def distinct_url_norms(self) -> set[str]:
        return {t.normalized for t in self.url_targets if t.normalized}

    @property
    def distinct_url_originals(self) -> set[str]:
        return {t.original for t in self.url_targets if t.original}


def _find_col(headers: list[str], wanted: str) -> int:
    def clean(v: str) -> str:
        return (v or "").strip().strip("\"'").lower()

    target = clean(wanted)
    for i, h in enumerate(headers):
        if clean(h) == target:
            return i
    raise SheetFormatError(
        f"Column {wanted!r} not found. Available columns: {[h for h in headers if h] or '<none>'}"
    )


def _iter_rows(sheet):
    for row in sheet.iter_rows(values_only=True):
        if row is None:
            continue
        yield row


def read_excel(cfg: Config) -> ExcelData:
    wb = load_workbook(cfg.xlsx_path, read_only=True, data_only=True)
    try:
        data = ExcelData()
        # --- URL sheet -----------------------------------------------------
        if cfg.sheet_urls not in wb.sheetnames:
            raise SheetFormatError(
                f"Sheet {cfg.sheet_urls!r} missing. Have: {wb.sheetnames}"
            )
        ws_urls = wb[cfg.sheet_urls]
        headers = None
        cols: dict[str, int] = {}
        for row in _iter_rows(ws_urls):
            if headers is None:
                headers = [c if c is not None else "" for c in row]
                cols["example"] = _find_col(headers, cfg.col_example_url)
                if cfg.include_canonical_urls:
                    cols["canonical"] = _find_col(headers, cfg.col_canonical_url)
                if cfg.col_section in {h.strip() for h in headers}:
                    cols["section"] = headers.index(next(h for h in headers if h.strip() == cfg.col_section.strip()))
                continue
            data.url_rows_total += 1
            cells = [c if c is not None else "" for c in row]
            section = cells[cols["section"]].strip() if "section" in cols and cols["section"] < len(cells) else ""
            for source, col_key in (
                (cfg.col_example_url, "example"),
                (cfg.col_canonical_url, "canonical"),
            ):
                if col_key not in cols:
                    continue
                idx = cols[col_key]
                raw = cells[idx].strip() if idx < len(cells) else ""
                if not raw:
                    data.url_rows_skipped += 1
                    continue
                data.url_targets.append(
                    UrlTarget(
                        source_column=source,
                        original=raw,
                        normalized=normalize_url(raw, cfg),
                        section=section,
                    )
                )

        # --- File sheet ----------------------------------------------------
        if cfg.sheet_docs not in wb.sheetnames:
            raise SheetFormatError(
                f"Sheet {cfg.sheet_docs!r} missing. Have: {wb.sheetnames}"
            )
        ws_docs = wb[cfg.sheet_docs]
        headers = None
        cols = {}
        for row in _iter_rows(ws_docs):
            if headers is None:
                headers = [c if c is not None else "" for c in row]
                cols["file"] = _find_col(headers, cfg.col_file_name)
                if cfg.col_section in {h.strip() for h in headers}:
                    cols["section"] = headers.index(next(h for h in headers if h.strip() == cfg.col_section.strip()))
                continue
            data.file_rows_total += 1
            cells = [c if c is not None else "" for c in row]
            section = cells[cols["section"]].strip() if "section" in cols and cols["section"] < len(cells) else ""
            raw = cells[cols["file"]].strip() if cols["file"] < len(cells) else ""
            if not raw:
                data.file_rows_skipped += 1
                continue
            data.file_targets.append(FileTarget(original=raw, section=section))
        return data
    finally:
        wb.close()