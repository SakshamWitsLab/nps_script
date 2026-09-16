#!/usr/bin/env python3
"""NPS DROP cleanup - delete URLs / documents listed in the triage workbook.

Usage:
    python main.py                        # dry-run by default (report only)
    python main.py --apply                # actually delete (asks y/N first)
    python main.py --env .env --apply
    python main.py --list-sheets          # print workbook sheet names and exit
"""

import argparse
import logging
import sys

from openpyxl import load_workbook

import config as config_mod
import db as db_mod
from chunk_plan import build_chunk_plan
from config import Config
from deleter import delete
from excel_reader import read_excel, read_keep
from logger import SectionLog, setup_logging
from matcher import Matcher
import reporter


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="NPS DROP cleanup tool")
    p.add_argument("--env", default=None, help="path to a .env file")
    p.add_argument("--apply", action="store_true",
                   help="run the real deletion (overrides DRY_RUN) - still asks y/N")
    p.add_argument("--dry-run", action="store_true",
                   help="force DRY_RUN=true (never touches the database)")
    p.add_argument("--delete-mode", choices=["hard", "soft"], default=None,
                   help="override DELETE_MODE for this run")
    p.add_argument("--show-config", action="store_true",
                   help="print the active configuration and exit")
    p.add_argument("--list-sheets", action="store_true",
                   help="print the workbook sheet names and exit")
    return p


def _print_detection_summary(log, excel, res, plan=None) -> None:
    breakdown = res.status_breakdown()
    url_example = sum(1 for t in excel.url_targets if not t.is_canonical)
    url_canonical = res.url_targets_used - url_example
    log.info("")
    log.info("Detection summary (from DROP sheets -> database):")
    log.info("  URLs   sheet rows: %d | example-url targets: %d | canonical targets: %d | "
             "distinct normalized: %d",
             excel.url_rows_total, url_example, url_canonical, res.url_normalized_targets)
    log.info("  Files  sheet rows: %d | file targets: %d",
             excel.file_rows_total, res.file_targets_used)
    log.info("")
    log.info("  URL documents matched in DB    : %d", len(set(res.url_ids)))
    log.info("  File documents matched in DB   : %d", len(set(res.file_ids)))
    log.info("  Total rows to remove (expanded): %d (%d via supersedes/children)",
             len(res.matched_ids), len(res.expansion_ids))
    total_chunks = sum(s["total"] for s in res.chunk_stats.values())
    active_chunks = sum(s["active"] for s in res.chunk_stats.values())
    log.info(
        "  Chunks attached                 : %d total / %d active", total_chunks, active_chunks
    )
    if breakdown["by_status"]:
        statuses = ", ".join(f"{k}={v}" for k, v in sorted(breakdown["by_status"].items()))
        log.info("  Matched status mix              : %s", statuses)
    if plan is not None:
        log.info("")
        log.info("Chunk plan (identity = document_chunks.doc_name):")
        for line in plan.summary_lines():
            log.info("%s", line)
    if res.url_misses or res.file_misses:
        log.info("")
        log.info("  Targets with NO DB match:")
        log.info("    - URLs : %d", len(res.url_misses))
        log.info("    - Files: %d", len(res.file_misses))


def _confirm_delete(cfg: Config, count: int, plan=None) -> bool:
    print()
    print(f"About to {cfg.delete_mode}-delete {count} document row(s).")
    if plan is not None:
        print(f"  - {plan.drop_chunks} chunk(s) matched by doc_name will be removed")
        print(f"  - {plan.survivors} surviving chunk(s) preserved "
              f"({len(plan.keep_chunk_ids)} KEEP, {len(plan.untouched_chunk_ids)} other)")
        print(f"  - {len(plan.reparent)} surviving chunk(s) reparented off DROP containers")
        if plan.anomalies:
            print(f"  - WARNING: {len(plan.anomalies)} chunk(s) have no surviving owner")
    print("  hard -> DELETE the planned chunks, then DELETE FROM documents")
    print("  soft -> chunks matching DROP marked is_deleted; documents status=TRASHED")
    answer = input("Proceed? [y/N] ").strip().lower()
    return answer in ("y", "yes")


def list_sheets(path: str) -> None:
    wb = load_workbook(path, read_only=True)
    try:
        print(f"Sheets in {path}:")
        for name in wb.sheetnames:
            ws = wb[name]
            print(f"  - {name!r} (dim {ws.max_row}x{ws.max_column})")
    finally:
        wb.close()


def main(argv=None) -> int:
    args = build_parser().parse_args(argv)

    cfg = config_mod.load_config(args.env)
    if args.dry_run:
        cfg.dry_run = True
    if args.apply:
        cfg.dry_run = False
    if args.delete_mode:
        cfg.delete_mode = args.delete_mode

    if args.list_sheets:
        list_sheets(cfg.xlsx_path)
        return 0

    if args.show_config:
        print(config_mod.print_summary(cfg))
        return 0

    errors = cfg.validate()
    if errors:
        print("Configuration errors:")
        for e in errors:
            print(f"  - {e}")
        print("Fix the .env / environment and retry. See .env.example.")
        return 2

    out_dir = reporter.make_output_dir(cfg)
    log: logging.Logger = setup_logging(cfg, log_dir=out_dir)
    section = SectionLog(log)

    log.info("NPS DROP cleanup starting (host=%s db=%s)",
             cfg.db_host or "(from DATABASE_URL)", cfg.db_name or "(from DATABASE_URL)")
    for line in config_mod.print_summary(cfg).splitlines():
        log.info("  %s", line)

    try:
        # STEP 1: parse workbook
        section.step("Read triage workbook",
                     f"{cfg.xlsx_path} -> sheets {cfg.sheet_urls!r} / {cfg.sheet_docs!r}")
        excel = read_excel(cfg)
        log.info("Parsed %d URL targets and %d file targets (skipped rows: urls=%d files=%d, "
                 "header-artifact rows ignored=%d).",
                 len(excel.url_targets), len(excel.file_targets),
                 excel.url_rows_skipped, excel.file_rows_skipped,
                 excel.header_artifact_rows)
        if not excel.url_targets and not excel.file_targets:
            log.error("No targets could be read from the workbook. Aborting.")
            return 3
        keep = read_keep(cfg)
        log.info("Parsed %d KEEP URL targets and %d KEEP file targets (used to "
                 "protect chunks).", len(keep.url_targets), len(keep.file_targets))

        # STEP 2: connect + statistics
        section.step("Connect to PostgreSQL")
        conn = db_mod.connect(cfg)
        try:
            stats_before = db_mod.table_stats(conn)
            log.info("DB snapshot: %s documents (%s URLs / %s files / %s soft-deleted), "
                     "%s chunks (%s active)",
                     stats_before["documents_total"]["documents"],
                     stats_before["documents_total"]["urls"],
                     stats_before["documents_total"]["docs"],
                     stats_before["documents_total"]["soft_deleted"],
                     stats_before["chunks_total"]["chunks"],
                     stats_before["chunks_total"]["active_chunks"])

            # STEP 3: detect
            section.step("Detect matches (URLs via normalized JOIN, files via strategy)",
                         f"delete_mode={cfg.delete_mode}, batch_size={cfg.batch_size}")
            matcher = Matcher(conn, cfg)
            res = matcher.detect(excel, log=log)

            # STEP 3b: plan chunk fate by identity (protects KEEP chunks)
            plan = build_chunk_plan(conn, cfg, excel, keep, res.matched_ids, log=log)
            _print_detection_summary(log, excel, res, plan)

            # STEP 4: pre reports
            section.step("Write pre-delete reports")
            pre_md, pre_json = reporter.write_pre(cfg, excel, res, out_dir, plan=plan)
            audit = reporter.write_matched_audit_csv(cfg, res, out_dir)
            log.info("Pre reports written: %s", ", ".join(str(x) for x in (pre_md, pre_json, audit)))

            # STEP 5: dry-run gate
            if cfg.dry_run:
                log.info("")
                log.info("DRY RUN ONLY - no rows were changed.")
                log.info("Review the pre_report.md above, then run with --apply "
                         "(or DRY_RUN=false) to actually delete.")
                conn.rollback()
                return 0

            matches = len(res.matched_ids)
            if matches == 0:
                log.info("Nothing to delete.")
                conn.rollback()
                return 0

            # STEP 6: confirm
            if not _confirm_delete(cfg, matches, plan):
                log.info("Deletion cancelled by user.")
                conn.rollback()
                return 0

            # STEP 7: delete
            section.step(f"Apply {cfg.delete_mode} deletion")
            del_stats = delete(conn, cfg, plan, log=log)
            log.info("Deleted %d documents / %d chunks (reparented %d) in %d batch(es) [%.1fs].",
                     del_stats.documents_affected, del_stats.chunks_affected,
                     del_stats.chunks_reparented, del_stats.batches, del_stats.duration_s)

            # STEP 8: post check + reports
            section.step("Post-delete verification", "re-running the same matching queries")
            post = matcher.post_check(excel)
            stats_after = db_mod.table_stats(conn)
            post_md, post_json = reporter.write_post(
                cfg, excel, post, del_stats, stats_after, out_dir
            )
            remain = reporter.write_remaining_csv(cfg, excel, post, out_dir)

            log.info("Post check: URL rows active remaining=%d, file rows active remaining=%d, "
                     "active chunk leftovers=%d",
                     len(post.url_active_remaining), len(post.file_active_remaining),
                     len(post.chunk_leftovers))
            if post.ok:
                log.info("VERIFICATION PASSED - nothing from the DROP list remains retrievable.")
            else:
                log.warning("VERIFICATION FAILED - leftovers detected. Details:")
                if post.url_active_remaining:
                    log.warning("  - %d URL document(s) still active", len(post.url_active_remaining))
                if post.file_active_remaining:
                    log.warning("  - %d file document(s) still active", len(post.file_active_remaining))
                if post.chunk_leftovers:
                    log.warning("  - %d active chunk(s) still carry a dropped name",
                                len(post.chunk_leftovers))
            log.info("Reports: %s", ", ".join(str(x) for x in (post_md, post_json, remain)))
            conn.commit()
        finally:
            conn.close()
    except KeyboardInterrupt:
        log.warning("Interrupted - no changes committed.")
        return 130
    except Exception:
        log.exception("Pipeline failed - no changes committed.")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())